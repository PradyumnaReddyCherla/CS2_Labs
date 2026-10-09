"""Integrity of the classroom sample."""
import json, yaml
from openpyxl import load_workbook
from src import sample_data as sd

def test_answer_key_blocks_exist():
    ids = {r["block_id"] for r in sd.records("deal_a") if r["doc_id"] == "DR-deal_a-001"}
    assert all(t["claim_block"] in ids for t in yaml.safe_load(open("data/dataroom/deal_a/answer_key.yaml"))["traps"])

def test_xlsx_has_literal_values():
    wb = load_workbook("data/dataroom/deal_a/raw/financials.xlsx", data_only=True)
    assert all(c.value is not None for ws in wb for row in ws.iter_rows(min_row=2) for c in row[1:])

def test_canaries_stay_in_their_deal():
    can = json.load(open("data/dataroom/canaries_visible.json"))
    a = " ".join(r["text"] for r in sd.records("deal_a")); b = " ".join(r["text"] for r in sd.records("deal_b"))
    assert all(x in a and x not in b for x in can["deal_a"]) and all(x in b and x not in a for x in can["deal_b"])

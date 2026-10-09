"""Shared, deterministic helpers for the classroom sample. Used by the notebooks and by
tools/make_fixtures.py, so prompts, fact ids and data splits always agree."""
import hashlib, json
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook

S = Path("data")
XLSX_DOC = "DR-deal_a-004"                      # financials.xlsx
ROW_MAP = {"Revenue": "revenue", "Gross profit": "gross_profit", "EBITDA": "ebitda", "Adjusted EBITDA": "adjusted_ebitda",
           "ARR (USD millions)": "arr", "Net revenue retention": "nrr", "Gross logo churn": "gross_logo_churn"}
SCALE = {"P&L": 1e6, "KPIs": None}            # P&L in USD millions; KPI rows scale per row below
KPI_SCALE = {"arr": 1e6, "nrr": 1.0, "gross_logo_churn": 1.0}
FEATURES = ["revenue_usd_m", "growth", "ebitda_margin", "recurring_share", "top_customer_share", "asking_ev_ebitda"]

def fact_id(deal_id, metric, period, doc_id):
    return "F-" + hashlib.sha1(f"{deal_id}|{metric}|{period}|{doc_id}".encode()).hexdigest()[:10]

def records(deal):
    return [json.loads(l) for l in open(S / f"dataroom/{deal}/export/records.jsonl")]

def seller_facts():
    """Tier-3 facts from financials.xlsx, values in absolute units, plus derived revenue growth."""
    wb = load_workbook(S / "dataroom/deal_a/raw/financials.xlsx", data_only=True)
    out = []
    for sheet in ("P&L", "KPIs"):
        rows = list(wb[sheet].iter_rows())
        header = [c.value for c in rows[0]]
        for row in rows[1:]:
            label = row[0].value
            if label not in ROW_MAP: continue
            m = ROW_MAP[label]
            for j, cell in enumerate(row[1:], start=1):
                scale = SCALE[sheet] or KPI_SCALE[m]
                out.append({"fact_id": fact_id("deal_a", m, header[j], XLSX_DOC), "deal_id": "deal_a", "metric": m, "period": header[j],
                            "value": float(cell.value) * scale, "unit": "USD" if scale == 1e6 else "pct", "tier": 3, "period_basis": "fiscal",
                            "citation": {"kind": "sheet_cell", "doc_id": XLSX_DOC, "cell": f"{sheet}!{cell.coordinate}"}})
    cu = wb["Customers"]; top = cu["B2"].value
    out.append({"fact_id": fact_id("deal_a", "top1_customer_share", "FY2025", XLSX_DOC), "deal_id": "deal_a", "metric": "top1_customer_share", "period": "FY2025",
                "value": float(top), "unit": "pct", "tier": 3, "period_basis": "fiscal", "citation": {"kind": "sheet_cell", "doc_id": XLSX_DOC, "cell": "Customers!B2"}})
    rev = {f["period"]: f for f in out if f["metric"] == "revenue"}
    g = (rev["FY2025"]["value"] - rev["FY2024"]["value"]) / rev["FY2024"]["value"]
    out.append({"fact_id": fact_id("deal_a", "revenue_growth", "FY2025", "derived"), "deal_id": "deal_a", "metric": "revenue_growth", "period": "FY2025",
                "value": round(g, 6), "unit": "pct", "tier": 3, "period_basis": "fiscal",
                "citation": {"kind": "derived", "doc_id": XLSX_DOC, "parents": [rev["FY2025"]["fact_id"], rev["FY2024"]["fact_id"]]}})
    return out

def claim_id(doc_id, block_id, i):
    return f"CL-{doc_id}-{block_id}-{i}"

def deal_log():
    return pd.read_csv(S / "reference/deal_log.csv")

def test_split():
    """The same split as the calibration slide (Lab 21)."""
    from sklearn.model_selection import train_test_split
    log = deal_log()
    return train_test_split(log[FEATURES], log["advanced_to_ic"], log["p_true_advanced_to_ic"],
                            test_size=0.3, random_state=7, stratify=log["advanced_to_ic"])

TASK_EVIDENCE = {"dev-w-01": "w-01", "dev-w-02": "w-05", "dev-w-04": "w-07"}   # dev task -> qrel with its gold evidence
CLAIM_BLOCKS = ["p0002_b001", "p0002_b002", "p0002_b003", "p0003_b001", "p0003_b002"]

def claim_rows(extracted):
    """extracted: {block_id: MgmtClaims}. Assign claim ids and align periods (CY2025 -> FY2025, basis calendar)."""
    rows = []
    for block_id in CLAIM_BLOCKS:
        for i, c in enumerate(extracted[block_id].claims):
            label = c.period
            basis = "calendar" if label.startswith("CY") else "fiscal"
            period = "FY" + label[2:] if basis == "calendar" else label
            rows.append({"claim_id": claim_id("DR-deal_a-001", block_id, i), "doc_id": "DR-deal_a-001", "block_id": block_id,
                         "metric": c.metric, "value": c.value, "scale": c.scale, "value_abs": c.value * c.scale, "unit": c.unit,
                         "period": period, "period_label": label, "period_basis": basis, "qualifier": c.qualifier, "quote": c.quote})
    return rows

def evidence_for(task_id):
    qrels = {q["qid"]: q for q in (json.loads(l) for l in open(S / "eval/qrels.jsonl"))}
    by_id = {f'{r["doc_id"]}:{r["block_id"]}': r for r in records("deal_a")}
    return [by_id[k] for k in qrels[TASK_EVIDENCE[task_id]]["relevant"]]

def dev_tasks():
    return {t["task_id"]: t for t in (json.loads(l) for l in open("evals/tasks/dev.jsonl"))}

def cim_text():
    return "\n\n".join(r["text"] for r in records("deal_a") if r["doc_id"] == "DR-deal_a-001")
